from pathlib import Path

from pydantic import field_validator

from ..models import StrictModel
from ..source_types import SourceCommand, SourceContext


class XlsxExtractTextInput(StrictModel):
    """Input for xlsx.extract_text.

    Attributes:
        path: Workspace-relative XLSX or XLSM path.
    """

    path: str

    @field_validator("path")
    @classmethod
    def require_workbook_extension(cls, value: str) -> str:
        """Requires the path to target an XLSX/XLSM workbook.

        Args:
            value: Workspace-relative path from verifier.json.

        Returns:
            The unchanged path when it has a supported workbook extension.

        Raises:
            ValueError: If the path does not end in ``.xlsx`` or ``.xlsm``.
        """
        if Path(value).suffix.lower() not in {".xlsx", ".xlsm"}:
            raise ValueError("xlsx.extract_text requires a .xlsx or .xlsm file")
        return value


class XlsxExtractTextOutput(StrictModel):
    """Output from xlsx.extract_text.

    Attributes:
        text: Extracted sheet names and non-empty cell rows.
    """

    text: str


class ExtractText(SourceCommand):
    """Extracts readable text from workbook sheets."""

    name = "extract_text"
    input_model = XlsxExtractTextInput
    output_model = XlsxExtractTextOutput

    def run(self, source_input: XlsxExtractTextInput, context: SourceContext) -> XlsxExtractTextOutput:
        """Runs XLSX text extraction.

        Args:
            source_input: Validated workbook extraction input.
            context: Source runtime context.

        Returns:
            Extracted workbook text.

        """
        resolved = context.resolve_path(source_input.path)
        return XlsxExtractTextOutput(text=extract_xlsx(resolved, context.max_content_chars))


def extract_xlsx(path: Path, limit: int) -> str:
    """Extracts sheet names and non-empty cell rows from a workbook.

    Args:
        path: Resolved workbook path.
        limit: Maximum number of characters to return.

    Returns:
        Extracted workbook text capped to ``limit`` characters.
    """
    from openpyxl import load_workbook

    wb = load_workbook(path, data_only=True, read_only=True)
    parts: list[str] = []
    total = 0
    try:
        for sheet_name in wb.sheetnames:
            total = append_limited(parts, f"[Sheet: {sheet_name}]", total, limit)
            if total >= limit:
                return "\n".join(parts)[:limit]
            for row in wb[sheet_name].iter_rows(values_only=True):
                cells = [str(cell) if cell is not None else "" for cell in row]
                if any(cell.strip() for cell in cells):
                    total = append_limited(parts, " | ".join(cells), total, limit)
                    if total >= limit:
                        return "\n".join(parts)[:limit]
        return "\n".join(parts)[:limit]
    finally:
        wb.close()


def append_limited(parts: list[str], text: str, total: int, limit: int) -> int:
    """Appends text and returns the updated character count.

    Args:
        parts: Text fragments collected so far.
        text: Text fragment to append.
        total: Current approximate character count.
        limit: Maximum desired character count.

    Returns:
        Updated approximate character count.
    """
    parts.append(text)
    return total + len(text) + 1


COMMANDS = (ExtractText(),)
